#!/usr/bin/env python3
"""Generate an advanced multi-doctor P&L workbook with consolidation + dashboard.

Input CSV columns:
- doctor
- month (YYYY-MM)
- patient_revenue
- procedure_revenue
- ancillary_revenue
- salary_expense
- supplies_expense
- facility_expense
- admin_expense
- malpractice_expense
- other_expense

Example:
python scripts/generate_pnl_report.py --input data/doctor_financials_sample.csv --output output/Doctor_PnL_Report.xlsx
"""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MONTHS = [
    "Jan",
    "Feb",
    "Mar",
    "Apr",
    "May",
    "Jun",
    "Jul",
    "Aug",
    "Sep",
    "Oct",
    "Nov",
    "Dec",
]

REVENUE_ITEMS = ["Patient Revenue", "Procedure Revenue", "Ancillary Revenue"]
EXPENSE_ITEMS = [
    "Salary Expense",
    "Supplies Expense",
    "Facility Expense",
    "Admin Expense",
    "Malpractice Expense",
    "Other Expense",
]

CSV_TO_ITEM = {
    "patient_revenue": "Patient Revenue",
    "procedure_revenue": "Procedure Revenue",
    "ancillary_revenue": "Ancillary Revenue",
    "salary_expense": "Salary Expense",
    "supplies_expense": "Supplies Expense",
    "facility_expense": "Facility Expense",
    "admin_expense": "Admin Expense",
    "malpractice_expense": "Malpractice Expense",
    "other_expense": "Other Expense",
}


@dataclass
class DoctorData:
    revenue: Dict[str, List[float]]
    expense: Dict[str, List[float]]


def parse_month_to_index(month_value: str) -> int:
    parsed = datetime.strptime(month_value, "%Y-%m")
    return parsed.month - 1


def to_float(value: str) -> float:
    try:
        return float(value)
    except (ValueError, TypeError):
        return 0.0


def load_financial_data(csv_path: Path) -> Dict[str, DoctorData]:
    doctor_rows = defaultdict(
        lambda: {
            "revenue": {item: [0.0] * 12 for item in REVENUE_ITEMS},
            "expense": {item: [0.0] * 12 for item in EXPENSE_ITEMS},
        }
    )

    with csv_path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        required_columns = {"doctor", "month", *CSV_TO_ITEM.keys()}
        missing = required_columns - set(reader.fieldnames or [])
        if missing:
            missing_cols = ", ".join(sorted(missing))
            raise ValueError(f"Missing required CSV columns: {missing_cols}")

        for row in reader:
            doctor = row["doctor"].strip()
            month_idx = parse_month_to_index(row["month"].strip())

            for csv_col, item_name in CSV_TO_ITEM.items():
                value = to_float(row.get(csv_col, 0))
                if item_name in REVENUE_ITEMS:
                    doctor_rows[doctor]["revenue"][item_name][month_idx] += value
                else:
                    doctor_rows[doctor]["expense"][item_name][month_idx] += value

    return {
        doctor: DoctorData(revenue=data["revenue"], expense=data["expense"])
        for doctor, data in sorted(doctor_rows.items())
    }


def style_headers(ws, row: int, start_col: int, end_col: int, fill_color: str) -> None:
    thin = Side(style="thin", color="D9D9D9")
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def apply_currency_row(ws, row: int, start_col: int = 2, end_col: int = 14) -> None:
    for col in range(start_col, end_col + 1):
        ws.cell(row=row, column=col).number_format = '$#,##0'


def create_doctor_sheet(wb: Workbook, doctor: str, data: DoctorData) -> None:
    ws = wb.create_sheet(title=doctor[:31])
    ws["A1"] = f"{doctor} - Profit & Loss"
    ws["A1"].font = Font(size=14, bold=True)

    ws["A3"] = "Revenue"
    ws["A3"].font = Font(bold=True, color="1F4E78")

    ws["A9"] = "Expenses"
    ws["A9"].font = Font(bold=True, color="9C0006")

    ws["A16"] = "Key Metrics"
    ws["A16"].font = Font(bold=True, color="385723")

    # Header row
    ws["A4"] = "Category"
    for idx, month in enumerate(MONTHS, start=2):
        ws.cell(row=4, column=idx, value=month)
    ws["N4"] = "Annual"
    style_headers(ws, 4, 1, 14, "1F4E78")

    # Revenue rows
    row = 5
    for item in REVENUE_ITEMS:
        ws.cell(row=row, column=1, value=item)
        for m in range(12):
            ws.cell(row=row, column=m + 2, value=data.revenue[item][m])
        ws.cell(row=row, column=14, value=f"=SUM(B{row}:M{row})")
        apply_currency_row(ws, row)
        row += 1

    ws.cell(row=8, column=1, value="Total Revenue").font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m + 2)
        ws.cell(row=8, column=m + 2, value=f"=SUM({col}5:{col}7)")
    ws.cell(row=8, column=14, value="=SUM(N5:N7)")
    apply_currency_row(ws, 8)

    # Expense rows
    row = 10
    for item in EXPENSE_ITEMS:
        ws.cell(row=row, column=1, value=item)
        for m in range(12):
            ws.cell(row=row, column=m + 2, value=data.expense[item][m])
        ws.cell(row=row, column=14, value=f"=SUM(B{row}:M{row})")
        apply_currency_row(ws, row)
        row += 1

    ws.cell(row=16, column=1, value="Total Expense").font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m + 2)
        ws.cell(row=16, column=m + 2, value=f"=SUM({col}10:{col}15)")
    ws.cell(row=16, column=14, value="=SUM(N10:N15)")
    apply_currency_row(ws, 16)

    ws.cell(row=17, column=1, value="Operating Profit").font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m + 2)
        ws.cell(row=17, column=m + 2, value=f"={col}8-{col}16")
    ws.cell(row=17, column=14, value="=N8-N16")
    apply_currency_row(ws, 17)

    ws.cell(row=18, column=1, value="Operating Margin %").font = Font(bold=True)
    for m in range(12):
        col = get_column_letter(m + 2)
        ws.cell(row=18, column=m + 2, value=f"=IFERROR({col}17/{col}8,0)")
        ws.cell(row=18, column=m + 2).number_format = "0.0%"
    ws.cell(row=18, column=14, value="=IFERROR(N17/N8,0)")
    ws.cell(row=18, column=14).number_format = "0.0%"

    ws.cell(row=20, column=1, value="Quick Insight").font = Font(bold=True)
    ws.cell(
        row=20,
        column=2,
        value='=IF(N17>0,"Profitable year","Needs cost optimization")',
    )

    for col in range(1, 15):
        width = 20 if col == 1 else 12
        ws.column_dimensions[get_column_letter(col)].width = width


def create_consolidated_sheet(wb: Workbook, doctors: List[str]) -> None:
    ws = wb.create_sheet(title="Consolidated")
    ws["A1"] = "Consolidated Profit & Loss"
    ws["A1"].font = Font(size=14, bold=True)

    ws["A3"] = "Category"
    for idx, month in enumerate(MONTHS, start=2):
        ws.cell(row=3, column=idx, value=month)
    ws["N3"] = "Annual"
    style_headers(ws, 3, 1, 14, "2F5597")

    labels = [
        "Total Revenue",
        "Total Expense",
        "Operating Profit",
        "Operating Margin %",
    ]

    for r, label in enumerate(labels, start=4):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)

    for m in range(12):
        col = get_column_letter(m + 2)
        ws.cell(
            row=4,
            column=m + 2,
            value="=" + "+".join([f"'{doc[:31]}'!{col}8" for doc in doctors]),
        )
        ws.cell(
            row=5,
            column=m + 2,
            value="=" + "+".join([f"'{doc[:31]}'!{col}16" for doc in doctors]),
        )
        ws.cell(row=6, column=m + 2, value=f"={col}4-{col}5")
        ws.cell(row=7, column=m + 2, value=f"=IFERROR({col}6/{col}4,0)")
        ws.cell(row=7, column=m + 2).number_format = "0.0%"

    ws["N4"] = "=SUM(B4:M4)"
    ws["N5"] = "=SUM(B5:M5)"
    ws["N6"] = "=N4-N5"
    ws["N7"] = "=IFERROR(N6/N4,0)"
    ws["N7"].number_format = "0.0%"

    apply_currency_row(ws, 4)
    apply_currency_row(ws, 5)
    apply_currency_row(ws, 6)

    ws["A10"] = "Doctor"
    ws["B10"] = "Annual Revenue"
    ws["C10"] = "Annual Expense"
    ws["D10"] = "Annual Profit"
    ws["E10"] = "Margin %"
    style_headers(ws, 10, 1, 5, "548235")

    for idx, doctor in enumerate(doctors, start=11):
        ws.cell(row=idx, column=1, value=doctor)
        ws.cell(row=idx, column=2, value=f"='{doctor[:31]}'!N8")
        ws.cell(row=idx, column=3, value=f"='{doctor[:31]}'!N16")
        ws.cell(row=idx, column=4, value=f"=B{idx}-C{idx}")
        ws.cell(row=idx, column=5, value=f"=IFERROR(D{idx}/B{idx},0)")
        ws.cell(row=idx, column=5).number_format = "0.0%"
        for c in (2, 3, 4):
            ws.cell(row=idx, column=c).number_format = '$#,##0'

    for col in range(1, 15):
        ws.column_dimensions[get_column_letter(col)].width = 18 if col == 1 else 12


def create_dashboard_sheet(wb: Workbook, doctor_count: int) -> None:
    ws = wb.create_sheet(title="Dashboard")
    ws["A1"] = "Executive Dashboard"
    ws["A1"].font = Font(size=16, bold=True)

    ws["A3"] = "Metric"
    ws["B3"] = "Value"
    style_headers(ws, 3, 1, 2, "7F6000")

    metrics = [
        ("Annual Revenue", "=Consolidated!N4", '$#,##0'),
        ("Annual Expense", "=Consolidated!N5", '$#,##0'),
        ("Annual Profit", "=Consolidated!N6", '$#,##0'),
        ("Operating Margin", "=Consolidated!N7", "0.0%"),
        ("Doctor Count", str(doctor_count), "0"),
    ]

    for idx, (metric, formula, fmt) in enumerate(metrics, start=4):
        ws.cell(row=idx, column=1, value=metric)
        ws.cell(row=idx, column=2, value=formula)
        ws.cell(row=idx, column=2).number_format = fmt

    # Monthly trend chart
    trend = LineChart()
    trend.title = "Monthly Revenue vs Expense"
    trend.style = 13
    trend.y_axis.title = "Amount ($)"
    trend.x_axis.title = "Month"

    data = Reference(wb["Consolidated"], min_col=2, max_col=13, min_row=4, max_row=5)
    cats = Reference(wb["Consolidated"], min_col=2, max_col=13, min_row=3, max_row=3)
    trend.add_data(data, titles_from_data=False, from_rows=True)
    trend.set_categories(cats)
    ws.add_chart(trend, "D3")

    # Profit by doctor chart
    bar = BarChart()
    bar.title = "Annual Profit by Doctor"
    bar.y_axis.title = "Profit ($)"
    bar.x_axis.title = "Doctor"

    start_row = 11
    end_row = 10 + doctor_count
    bar_data = Reference(wb["Consolidated"], min_col=4, min_row=10, max_row=end_row)
    bar_cats = Reference(wb["Consolidated"], min_col=1, min_row=start_row, max_row=end_row)
    bar.add_data(bar_data, titles_from_data=True)
    bar.set_categories(bar_cats)
    ws.add_chart(bar, "D20")

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 14


def build_report(input_csv: Path, output_xlsx: Path) -> None:
    doctor_data = load_financial_data(input_csv)
    doctors = list(doctor_data.keys())

    if len(doctors) != 18:
        raise ValueError(
            f"Expected exactly 18 doctors in input data, found {len(doctors)}."
        )

    wb = Workbook()
    wb.remove(wb.active)

    for doctor, data in doctor_data.items():
        create_doctor_sheet(wb, doctor, data)

    create_consolidated_sheet(wb, doctors)
    create_dashboard_sheet(wb, len(doctors))

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate a multi-doctor P&L workbook with consolidation dashboard."
    )
    parser.add_argument("--input", required=True, type=Path, help="Input CSV file")
    parser.add_argument("--output", required=True, type=Path, help="Output XLSX file")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build_report(args.input, args.output)
    print(f"Report generated: {args.output}")


if __name__ == "__main__":
    main()
